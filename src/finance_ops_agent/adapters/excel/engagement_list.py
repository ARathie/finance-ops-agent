"""Read Kevin's engagement list: engagements.xlsx, or a folder of CSV exports.

Only reads cells as text, with the row numbers Kevin sees in Excel; making
sense of them is domain code. The agent never writes to these files.
"""

import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from finance_ops_agent.domain.engagements import RawRow, RawWorkbook

SHEET_NAMES = ("Clients", "Consultants", "Vendors", "Engagements")


class MissingSheetError(Exception):
    """The workbook has no sheet with a required name."""


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        # Excel gives numeric cells as floats. Dollars-with-cents survive the
        # shortest-repr round trip exactly; the domain parses the text back to
        # integers, so no float ever reaches a calculation.
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def _sheet_rows(sheet: Worksheet) -> list[RawRow]:
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        return []
    headers = [_cell_text(cell) for cell in header_cells]
    rows: list[RawRow] = []
    for row_number, cells in enumerate(rows_iter, start=2):
        texts = [_cell_text(cell) for cell in cells]
        if not any(texts):
            continue
        rows.append(RawRow(row_number, dict(zip(headers, texts, strict=False))))
    return rows


class ExcelEngagementList:
    """Reads engagements.xlsx with the four sheets from docs/engagement-list.md."""

    def __init__(self, path: Path) -> None:
        self.path = path

    def load(self) -> RawWorkbook:
        workbook = load_workbook(self.path, read_only=False, data_only=True)
        sheets: dict[str, list[RawRow]] = {}
        for name in SHEET_NAMES:
            if name not in workbook.sheetnames:
                raise MissingSheetError(f'{self.path.name} has no "{name}" sheet')
            sheets[name] = _sheet_rows(workbook[name])
        return RawWorkbook(
            clients=sheets["Clients"],
            consultants=sheets["Consultants"],
            vendors=sheets["Vendors"],
            engagements=sheets["Engagements"],
        )


class CsvEngagementList:
    """Reads a folder of CSV exports: clients.csv, consultants.csv, vendors.csv,
    engagements.csv."""

    def __init__(self, directory: Path) -> None:
        self.directory = directory

    def _file_rows(self, name: str) -> list[RawRow]:
        path = self.directory / f"{name.lower()}.csv"
        if not path.exists():
            raise MissingSheetError(f"{self.directory} has no {path.name}")
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.reader(handle)
            try:
                headers = [cell.strip() for cell in next(reader)]
            except StopIteration:
                return []
            rows: list[RawRow] = []
            for row_number, cells in enumerate(reader, start=2):
                texts = [cell.strip() for cell in cells]
                if not any(texts):
                    continue
                rows.append(RawRow(row_number, dict(zip(headers, texts, strict=False))))
            return rows

    def load(self) -> RawWorkbook:
        return RawWorkbook(
            clients=self._file_rows("Clients"),
            consultants=self._file_rows("Consultants"),
            vendors=self._file_rows("Vendors"),
            engagements=self._file_rows("Engagements"),
        )
