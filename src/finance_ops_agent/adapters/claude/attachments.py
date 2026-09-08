"""Turn an attachment into content blocks Claude can read.

PDFs go as documents (32 MB / 600 page limits), images as images (very large
screenshots are downscaled first), spreadsheets and .docx as text; anything
else is CANT_READ_ATTACHMENT. See docs/integrations/claude-extraction.md.
"""

import base64
import csv
import io
import re
import zipfile
from pathlib import PurePosixPath

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from finance_ops_agent.ports.reader import CantReadAttachmentError

MAX_PDF_BYTES = 32 * 1024 * 1024
MAX_PDF_PAGES = 600
MAX_SHEET_ROWS = 500
MAX_IMAGE_SIDE = 2048

IMAGE_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}


def _pdf_blocks(content: bytes, filename: str) -> list[dict[str, object]]:
    if len(content) > MAX_PDF_BYTES:
        raise CantReadAttachmentError(f"{filename} is larger than 32 MB")
    try:
        pages = len(PdfReader(io.BytesIO(content)).pages)
    except Exception as error:
        raise CantReadAttachmentError(f"{filename} is not a readable PDF") from error
    if pages > MAX_PDF_PAGES:
        raise CantReadAttachmentError(f"{filename} has more than {MAX_PDF_PAGES} pages")
    return [
        {
            "type": "document",
            "source": {
                "type": "base64",
                "media_type": "application/pdf",
                "data": base64.standard_b64encode(content).decode("ascii"),
            },
        }
    ]


def _image_blocks(content: bytes, filename: str, media_type: str) -> list[dict[str, object]]:
    try:
        image = Image.open(io.BytesIO(content))
        image.load()
    except Exception as error:
        raise CantReadAttachmentError(f"{filename} is not a readable image") from error
    if max(image.size) > MAX_IMAGE_SIDE:
        image.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        content, media_type = buffer.getvalue(), "image/png"
    return [
        {
            "type": "image",
            "source": {
                "type": "base64",
                "media_type": media_type,
                "data": base64.standard_b64encode(content).decode("ascii"),
            },
        }
    ]


def _text_block(filename: str, text: str) -> list[dict[str, object]]:
    return [{"type": "text", "text": f"Contents of {filename}:\n\n{text}"}]


def _xlsx_text(content: bytes, filename: str) -> str:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise CantReadAttachmentError(f"{filename} is not a readable spreadsheet") from error
    parts: list[str] = []
    for sheet in workbook.worksheets:
        lines = [f"# Sheet: {sheet.title}"]
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_SHEET_ROWS:
                lines.append(f"... ({MAX_SHEET_ROWS} row limit reached)")
                break
            lines.append("\t".join("" if cell is None else str(cell) for cell in row))
        parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _csv_text(content: bytes, filename: str) -> str:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    lines: list[str] = []
    for index, row in enumerate(csv.reader(io.StringIO(text))):
        if index >= MAX_SHEET_ROWS:
            lines.append(f"... ({MAX_SHEET_ROWS} row limit reached)")
            break
        lines.append("\t".join(row))
    return "\n".join(lines)


def _docx_text(content: bytes, filename: str) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
    except Exception as error:
        raise CantReadAttachmentError(f"{filename} is not a readable document") from error
    # Paragraph breaks, then strip the remaining tags.
    xml = re.sub(r"</w:p>", "\n", xml)
    return re.sub(r"<[^>]+>", "", xml).strip()


def to_content_blocks(content: bytes, filename: str) -> list[dict[str, object]]:
    """Content blocks for the attachment, or CantReadAttachmentError."""
    suffix = PurePosixPath(filename.lower()).suffix
    if suffix == ".pdf":
        return _pdf_blocks(content, filename)
    if suffix in IMAGE_MEDIA_TYPES:
        return _image_blocks(content, filename, IMAGE_MEDIA_TYPES[suffix])
    if suffix == ".xlsx":
        return _text_block(filename, _xlsx_text(content, filename))
    if suffix == ".csv":
        return _text_block(filename, _csv_text(content, filename))
    if suffix == ".docx":
        return _text_block(filename, _docx_text(content, filename))
    raise CantReadAttachmentError(f"I can't read {suffix or 'this'} files ({filename})")
