"""The EngagementList contract: the fake, the .xlsx reader, and the CSV reader
all deliver the same raw rows for the same content, and the sample workbook
parses cleanly end to end.
"""

import csv
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from finance_ops_agent.adapters.excel.engagement_list import (
    CsvEngagementList,
    ExcelEngagementList,
    MissingSheetError,
)
from finance_ops_agent.adapters.fakes.engagement_list import FakeEngagementList
from finance_ops_agent.domain.engagements import RawWorkbook, parse_workbook
from finance_ops_agent.domain.money import Money
from finance_ops_agent.ports.engagement_list import EngagementList

CLIENT_HEADERS = [
    "Client",
    "Legal name",
    "Billing contact",
    "Billing email",
    "CC email",
    "Payment terms (days)",
    "Delivery",
    "Time system",
    "Names on timesheets",
    "Email domains",
    "QuickBooks customer",
    "Notes",
    "Active",
]
CONSULTANT_HEADERS = [
    "Consultant",
    "Other names",
    "Email",
    "Type",
    "Vendor company",
    "Paid by",
    "Pay timing (days)",
    "Active",
]
VENDOR_HEADERS = ["Vendor company", "Contact email", "Paid by", "Pay timing (days)", "Active"]
ENGAGEMENT_HEADERS = [
    "Consultant",
    "Client",
    "End client",
    "Role",
    "Start date",
    "End date",
    "Billing schedule",
    "First period start",
    "Bill rate",
    "Pay rate",
    "Rates from",
    "Send automatically",
    "Active",
]

CLIENT_VALUES = [
    "Acme Corp",
    "Acme Corporation",
    "Accounts Payable",
    "ap@acme.example",
    "",
    "30",
    "email",
    "Fieldglass",
    "Acme; ACME Corp.",
    "acme.example",
    "Acme Corporation",
    "",
    "yes",
]
CONSULTANT_VALUES = [
    "Priya Shah",
    "P. Shah",
    "priya@example.com",
    "contractor",
    "",
    "bank transfer",
    "15",
    "yes",
]
VENDOR_VALUES = ["Blue Peak Consulting LLC", "billing@bluepeak.example", "check", "30", "yes"]
ENGAGEMENT_VALUES = [
    "Priya Shah",
    "Acme Corp",
    "",
    "Senior PeopleSoft Developer",
    "2026-02-01",
    "",
    "monthly",
    "",
    "140.00",
    "100.00",
    "2026-02-01",
    "no",
    "yes",
]


def build_xlsx(path: Path) -> None:
    """The sample workbook, using the cell types Excel really produces:
    numbers for rates and terms, real dates for date columns."""
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]

    def add(name: str, headers: list[str], values: list[str]) -> None:
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        date_columns = {"Start date", "End date", "First period start", "Rates from"}
        typed: list[object] = []
        for header, value in zip(headers, values, strict=True):
            if header in date_columns:
                typed.append(date.fromisoformat(value) if value else "")
            elif header in ("Payment terms (days)", "Pay timing (days)"):
                typed.append(int(value))
            elif header in ("Bill rate", "Pay rate"):
                typed.append(float(value))
            else:
                typed.append(value)
        sheet.append(typed)
        sheet.append([""] * len(headers))  # a blank row, which readers must skip

    add("Clients", CLIENT_HEADERS, CLIENT_VALUES)
    add("Consultants", CONSULTANT_HEADERS, CONSULTANT_VALUES)
    add("Vendors", VENDOR_HEADERS, VENDOR_VALUES)
    add("Engagements", ENGAGEMENT_HEADERS, ENGAGEMENT_VALUES)
    workbook.save(path)


def build_csvs(directory: Path) -> None:
    for name, headers, values in (
        ("clients", CLIENT_HEADERS, CLIENT_VALUES),
        ("consultants", CONSULTANT_HEADERS, CONSULTANT_VALUES),
        ("vendors", VENDOR_HEADERS, VENDOR_VALUES),
        ("engagements", ENGAGEMENT_HEADERS, ENGAGEMENT_VALUES),
    ):
        with (directory / f"{name}.csv").open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(headers)
            writer.writerow(values)
            writer.writerow([""] * len(headers))


def fake_workbook() -> RawWorkbook:
    from finance_ops_agent.domain.engagements import RawRow

    def rows(headers: list[str], values: list[str]) -> list[RawRow]:
        return [RawRow(2, dict(zip(headers, values, strict=True)))]

    return RawWorkbook(
        clients=rows(CLIENT_HEADERS, CLIENT_VALUES),
        consultants=rows(CONSULTANT_HEADERS, CONSULTANT_VALUES),
        vendors=rows(VENDOR_HEADERS, VENDOR_VALUES),
        engagements=rows(ENGAGEMENT_HEADERS, ENGAGEMENT_VALUES),
    )


@pytest.fixture(params=["fake", "xlsx", "csv"])
def engagement_list(request: pytest.FixtureRequest, tmp_path: Path) -> EngagementList:
    if request.param == "fake":
        return FakeEngagementList(fake_workbook())
    if request.param == "xlsx":
        path = tmp_path / "engagements.xlsx"
        build_xlsx(path)
        return ExcelEngagementList(path)
    build_csvs(tmp_path)
    return CsvEngagementList(tmp_path)


def test_the_sample_workbook_loads_and_parses(engagement_list: EngagementList) -> None:
    raw = engagement_list.load()
    assert [row.row_number for row in raw.clients] == [2]
    parsed = parse_workbook(raw)
    assert parsed.problems == []
    assert parsed.clients[0].name == "Acme Corp"
    assert parsed.clients[0].payment_terms_days == 30
    assert parsed.consultants[0].emails == ("priya@example.com",)
    assert parsed.vendors[0].company == "Blue Peak Consulting LLC"
    engagement = parsed.engagements[0]
    assert engagement.bill_rate == Money(14_000)
    assert engagement.pay_rate == Money(10_000)
    assert engagement.start_date == date(2026, 2, 1)
    assert engagement.rates_from == date(2026, 2, 1)


def test_a_bad_row_is_reported_with_its_sheet_and_row(
    engagement_list: EngagementList, tmp_path: Path
) -> None:
    raw = engagement_list.load()
    broken = RawWorkbook(
        clients=raw.clients,
        consultants=raw.consultants,
        vendors=raw.vendors,
        engagements=[
            type(raw.engagements[0])(
                raw.engagements[0].row_number,
                {**raw.engagements[0].cells, "Bill rate": ""},
            )
        ],
    )
    parsed = parse_workbook(broken)
    assert parsed.problems
    assert parsed.problems[0].sheet == "Engagements"
    assert parsed.problems[0].row_number == 2


def test_missing_sheet_is_an_error(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "Clients"  # type: ignore[union-attr]
    path = tmp_path / "half.xlsx"
    workbook.save(path)
    with pytest.raises(MissingSheetError):
        ExcelEngagementList(path).load()

    with pytest.raises(MissingSheetError):
        CsvEngagementList(tmp_path / "empty").load()
