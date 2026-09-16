"""Build `templates/engagements-template.xlsx`, the starter engagement list.

Run `uv run python templates/build_template.py` after changing a column in
`docs/engagement-list.md`; `tests/unit/test_engagements_template.py` fails if
the committed workbook and the reader ever disagree.

The columns, their order, and the example values are the ones in
`docs/engagement-list.md`. Every example row is made up, and together they are
the worked example from that page: Priya Shah at Acme Corp, 156 hours at
140.00 billed and 100.00 paid.
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

OUTPUT = Path(__file__).resolve().parent / "engagements-template.xlsx"

HEADER_FILL = PatternFill("solid", start_color="FFE8EEF7")
DATE_FORMAT = "yyyy-mm-dd"
MONEY_FORMAT = "0.00"

# Columns that may only hold one of a short list of words. A dropdown here
# saves the commonest LIST_ROW_PROBLEM there is: a typo in one of these.
CHOICES = {
    "Delivery": ("email", "portal"),
    "Type": ("employee", "contractor", "vendor"),
    "Paid by": ("bank transfer", "payroll", "check", "other"),
    "Billing schedule": ("monthly", "twice a month", "every two weeks", "weekly"),
    "Active": ("yes", "no"),
    "Send automatically": ("yes", "no"),
}
# Vendors are never on Icon's payroll (domain/engagements.py forbids it).
VENDOR_PAID_BY = ("bank transfer", "check", "other")

SHEETS: dict[str, tuple[list[str], list[object]]] = {
    "Clients": (
        [
            "Client",
            "Legal name",
            "Billing contact",
            "Billing email",
            "CC email",
            "Payment terms (days)",
            "Delivery",
            "Time system",
            "Names on timesheets",
            "Email domains",
            "QuickBooks customer",
            "Notes",
            "Active",
        ],
        [
            "Acme Corp",
            "Acme Corporation",
            "Accounts Payable",
            "ap@acme.example",
            "jane.doe@acme.example",
            30,
            "email",
            "Fieldglass",
            "Acme; ACME Corp.",
            "acme.example",
            "Acme Corporation",
            "Example row - replace with a real client",
            "yes",
        ],
    ),
    "Consultants": (
        [
            "Consultant",
            "Other names",
            "Email",
            "Type",
            "Vendor company",
            "Paid by",
            "Pay timing (days)",
            "Active",
        ],
        [
            "Priya Shah",
            "P. Shah; Shah, Priya",
            "priya@example.com",
            "contractor",
            "",
            "bank transfer",
            15,
            "yes",
        ],
    ),
    "Vendors": (
        [
            "Vendor company",
            "Contact email",
            "Paid by",
            "Pay timing (days)",
            "Active",
        ],
        [
            "Blue Peak Consulting LLC",
            "billing@bluepeak.example",
            "bank transfer",
            30,
            "yes",
        ],
    ),
    "Engagements": (
        [
            "Consultant",
            "Client",
            "End client",
            "Role",
            "Start date",
            "End date",
            "Billing schedule",
            "First period start",
            "Bill rate",
            "Pay rate",
            "Rates from",
            "Send automatically",
            "Active",
        ],
        [
            "Priya Shah",
            "Acme Corp",
            "",
            "Senior PeopleSoft Developer",
            date(2026, 2, 1),
            "",
            "monthly",
            "",
            140.00,
            100.00,
            date(2026, 2, 1),
            "no",
            "yes",
        ],
    ),
}

READ_ME = [
    ("How to fill this in", None),
    ("", None),
    ("Every sheet has one made-up example row. Replace it with your own, or", None),
    ("delete it once you have added real rows. The agent reads the four sheets", None),
    ("named Clients, Consultants, Vendors, and Engagements, and ignores this one.", None),
    ("", None),
    ("Clients", "One row per company Icon sends an invoice to."),
    ("Consultants", "One row per person doing the work."),
    ("Vendors", "Only needed when a consultant's Type is 'vendor'. Otherwise leave it."),
    ("Engagements", "One row per consultant working for one client."),
    ("", None),
    ("Rates", "Bill rate is what the client pays Icon. Pay rate is what Icon pays"),
    ("", "the consultant or their vendor. They are different numbers and the"),
    ("", "agent takes both from this file and nowhere else."),
    ("", None),
    ("Changing a rate", "Add a new Engagements row with the new rates and the date they"),
    ("", "start in 'Rates from'. Leave the old row alone: it is the record of"),
    ("", "what was charged before."),
    ("", None),
    ("Dates", "Type them as 2026-02-01. Excel may show them differently once you"),
    ("", "press Enter; that is fine, the agent reads the underlying date."),
    ("", None),
    ("Money", "Dollars with cents, like 140.00. No currency symbol needed."),
    ("", None),
    ("Yes / no columns", "Pick from the dropdown. Blank is not the same as 'no'."),
    ("", None),
    ("If something is wrong", "The agent never guesses and never edits this file. It emails"),
    ("", "Kevin naming the sheet and the row number, and waits."),
]


def add_sheet(book: Workbook, name: str, headers: list[str], example: list[object]) -> Worksheet:
    sheet: Worksheet = book.create_sheet(name)
    sheet.append(headers)
    sheet.append(example)
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        width = max(len(header) + 2, len(str(example[index - 1])) + 2, 12)
        sheet.column_dimensions[letter].width = min(width, 34)
        choices = (
            VENDOR_PAID_BY if (name == "Vendors" and header == "Paid by") else CHOICES.get(header)
        )
        if choices:
            rule = DataValidation(
                type="list",
                formula1='"' + ",".join(choices) + '"',
                allow_blank=True,
                showDropDown=False,
            )
            sheet.add_data_validation(rule)
            rule.add(f"{letter}2:{letter}400")
        if "date" in header.lower() or header == "Rates from":
            for row in range(2, 401):
                sheet.cell(row=row, column=index).number_format = DATE_FORMAT
        if header in ("Bill rate", "Pay rate"):
            for row in range(2, 401):
                sheet.cell(row=row, column=index).number_format = MONEY_FORMAT
    sheet.freeze_panes = "A2"
    return sheet


def build() -> Workbook:
    book = Workbook()
    first = book.active
    if first is not None:
        book.remove(first)
    guide: Worksheet = book.create_sheet("Read me")
    for label, detail in READ_ME:
        guide.append([label, detail])
    guide["A1"].font = Font(bold=True, size=14)
    for row in range(2, len(READ_ME) + 2):
        guide.cell(row=row, column=1).font = Font(bold=True)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 72
    for name, (headers, example) in SHEETS.items():
        add_sheet(book, name, headers, example)
    return book


if __name__ == "__main__":
    build().save(OUTPUT)
    print(f"wrote {OUTPUT}")
